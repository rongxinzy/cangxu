"""Read-only workbook parsing helpers for staged Cangxu imports."""

import re
from datetime import date, datetime, time
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_WORKBOOK_BYTES = 25 * 1024 * 1024
MAX_WORKBOOK_CELLS = 500_000
MAX_WORKBOOK_ROWS = 50_000
PART_USAGE_HEADER = '单机用量'
DESIGNATOR_HEADER = '位号'
DESIGNATOR_SEPARATOR = re.compile(r'[,，]')

COMMON_HEADERS = {
    '数量': 'quantity',
    '数量（PCS）': 'quantity',  # noqa: RUF001 - exact source workbook header
    'SN': 'serial',
    'SN号': 'serial',
    'PN号': 'part_number',
    '料号': 'part_number',
    '本年度生产批次': 'batch_code',
    '生产日期': 'production_date',
    'BOM号': 'bom_code',
    '供应商': 'supplier_label',
}
IMPORT_HEADERS = {
    'opening_stock': {
        **COMMON_HEADERS,
        '品名': 'part_name',
        '产品名称': 'part_name',
        '产品型号': 'part_model',
        '类别': 'part_category',
        'OEM PN 号': 'oem_part_number',
        '是否到货': 'receipt_status',
    },
    'part_catalog': {
        **COMMON_HEADERS,
        '品名': 'part_name',
        '产品名称': 'part_name',
        '产品型号': 'part_model',
        '类别': 'part_category',
        'OEM PN 号': 'oem_part_number',
        '是否到货': 'receipt_status',
    },
    'purchase_history': {
        **COMMON_HEADERS,
        '品名': 'part_name',
        '产品名称': 'part_name',
        '产品型号': 'part_model',
        '类别': 'part_category',
        '是否到货': 'receipt_status',
    },
    'shipment_history': {
        '数量': 'quantity',
        '数量（PCS）': 'quantity',  # noqa: RUF001 - exact source workbook header
        '出库产品': 'product_description',
        '状态': 'shipment_status',
        '状态（已出库/未出库）': 'shipment_status',  # noqa: RUF001 - source header
        '出库时间': 'shipped_at',
        '流程编号': 'external_reference',
        '项目名称': 'project_reference',
    },
    'service_history': {
        '料号': 'part_number',
        '名称': 'part_name',
        '数量': 'quantity',
        'SN': 'serial',
        '故障原因具体描述': 'fault_description',
        '售后状态': 'service_status',
        '售后出库': 'shipped_at',
        '售后入库': 'returned_at',
        '原因': 'failure_reason',
        '备注': 'notes',
    },
}


class WorkbookImportError(ValueError):
    """Raised when an uploaded workbook cannot be safely staged."""


def _json_value(value):
    """Convert supported Excel values into stable JSON-compatible scalars."""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def parse_workbook_rows(upload):
    """Return populated worksheet rows while preserving cell coordinates/format."""
    if upload.size > MAX_WORKBOOK_BYTES:
        raise WorkbookImportError('workbook_too_large')

    try:
        workbook = load_workbook(upload, read_only=True, data_only=False)
    except Exception as exc:
        raise WorkbookImportError('invalid_xlsx') from exc

    rows = []
    cell_count = 0
    try:
        for sheet in workbook.worksheets:
            for row_index, row in enumerate(sheet.iter_rows(), start=1):
                cell_count += len(row)
                if cell_count > MAX_WORKBOOK_CELLS:
                    raise WorkbookImportError('workbook_too_many_cells')

                populated = [
                    (column_index, cell)
                    for column_index, cell in enumerate(row, start=1)
                    if cell.value is not None
                ]
                if not populated:
                    continue

                values = {
                    get_column_letter(column_index): {
                        'value': _json_value(cell.value),
                        'data_type': cell.data_type,
                        'number_format': cell.number_format,
                    }
                    for column_index, cell in enumerate(row, start=1)
                }
                first_col = min(column for column, _cell in populated)
                last_col = max(column for column, _cell in populated)
                rows.append({
                    'source_sheet': sheet.title[:100],
                    'source_row': row_index,
                    'source_range': (
                        f'{get_column_letter(first_col)}{row_index}:'
                        f'{get_column_letter(last_col)}{row_index}'
                    ),
                    'source_values': values,
                    'normalized_values': {},
                })
                if len(rows) > MAX_WORKBOOK_ROWS:
                    raise WorkbookImportError('workbook_too_many_rows')
    except WorkbookImportError:
        raise
    except Exception as exc:
        raise WorkbookImportError('invalid_xlsx') from exc
    finally:
        workbook.close()

    if not rows:
        raise WorkbookImportError('workbook_has_no_populated_rows')

    return rows


def validate_part_usage_workbook(upload, sheet_name=None):
    """Validate part usage counts and designators without persisting the workbook."""
    if upload.size > MAX_WORKBOOK_BYTES:
        raise WorkbookImportError('workbook_too_large')

    try:
        workbook = load_workbook(upload, read_only=True, data_only=False)
    except Exception as exc:
        raise WorkbookImportError('invalid_xlsx') from exc

    cell_count = 0
    try:
        sheets = [sheet.title for sheet in workbook.worksheets]
        if not sheets:
            raise WorkbookImportError('workbook_has_no_populated_rows')
        if len(sheets) > 1 and not sheet_name:
            return {
                'status': 'sheet_selection_required',
                'sheets': sheets,
                'sheet': None,
                'checked_rows': 0,
                'workbook_errors': [],
                'invalid_rows': [],
            }

        selected_sheet = sheet_name or sheets[0]
        if selected_sheet not in sheets:
            raise WorkbookImportError('worksheet_not_found')

        sheet = workbook[selected_sheet]

        header_row = next(
            sheet.iter_rows(min_row=1, max_row=1, values_only=False), ()
        )
        cell_count += len(header_row)
        if cell_count > MAX_WORKBOOK_CELLS:
            raise WorkbookImportError('workbook_too_many_cells')

        usage_column_index = 2
        designator_column_index = 3
        workbook_errors = []
        for field, expected, column, column_index in (
            ('single_usage', PART_USAGE_HEADER, 'B', usage_column_index),
            ('designators', DESIGNATOR_HEADER, 'C', designator_column_index),
        ):
            cell = _row_cell(header_row, column_index)
            actual = cell.value if cell is not None else None
            label = actual.strip() if isinstance(actual, str) else actual
            if label != expected:
                workbook_errors.append({
                    'code': 'invalid_header',
                    'field': field,
                    'column': column,
                    'expected': expected,
                    'actual': _json_value(actual),
                })

        if workbook_errors:
            return {
                'status': 'invalid',
                'sheets': sheets,
                'sheet': selected_sheet,
                'checked_rows': 0,
                'workbook_errors': workbook_errors,
                'invalid_rows': [],
            }

        checked_rows = 0
        invalid_rows = []

        for row_index, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=False), start=2
        ):
            cell_count += len(row)
            if cell_count > MAX_WORKBOOK_CELLS:
                raise WorkbookImportError('workbook_too_many_cells')

            usage_cell = _row_cell(row, usage_column_index)
            designator_cell = _row_cell(row, designator_column_index)
            usage_value = usage_cell.value if usage_cell is not None else None
            designator_value = (
                designator_cell.value if designator_cell is not None else None
            )
            if _is_blank(usage_value) and _is_blank(designator_value):
                continue

            row_values = {
                get_column_letter(column_index): _json_value(cell.value)
                for column_index, cell in enumerate(row, start=1)
                if not _is_blank(cell.value)
            }
            checked_rows += 1
            if checked_rows > MAX_WORKBOOK_ROWS:
                raise WorkbookImportError('workbook_too_many_rows')

            usage, usage_valid, usage_blank = _part_usage_quantity(usage_value)
            designators, empty_segments = _part_designators(designator_value)
            errors = []

            if not usage_blank and not usage_valid:
                errors.append({
                    'code': 'invalid_single_usage',
                    'value': _json_value(usage_value),
                })
            elif usage_blank and designators:
                errors.append({
                    'code': 'missing_single_usage',
                    'actual': len(designators),
                })

            if empty_segments:
                errors.append({
                    'code': 'invalid_designator_separator',
                    'segments': empty_segments,
                })
            elif not usage_blank and usage_valid and usage != len(designators):
                errors.append({
                    'code': 'designator_count_mismatch',
                    'expected': usage,
                    'actual': len(designators),
                })

            if errors:
                invalid_rows.append({
                    'row': row_index,
                    'single_usage': _json_value(usage_value),
                    'designators': (
                        '' if _is_blank(designator_value) else str(designator_value)
                    ),
                    'designator_count': len(designators),
                    'row_values': row_values,
                    'errors': errors,
                })

        if checked_rows == 0:
            return {
                'status': 'invalid',
                'sheets': sheets,
                'sheet': selected_sheet,
                'checked_rows': 0,
                'workbook_errors': [{'code': 'no_data_rows'}],
                'invalid_rows': [],
            }

        return {
            'status': 'invalid' if invalid_rows else 'valid',
            'sheets': sheets,
            'sheet': selected_sheet,
            'checked_rows': checked_rows,
            'workbook_errors': [],
            'invalid_rows': invalid_rows,
        }
    except WorkbookImportError:
        raise
    except Exception as exc:
        raise WorkbookImportError('invalid_xlsx') from exc
    finally:
        workbook.close()


def _row_cell(row, column_index):
    """Return a cell from a potentially shortened read-only row tuple."""
    if column_index > len(row):
        return None
    return row[column_index - 1]


def _is_blank(value):
    """Return whether an Excel value represents an empty cell."""
    return value is None or (isinstance(value, str) and not value.strip())


def _part_usage_quantity(value):
    """Return (quantity, is_valid, is_blank) for a part usage cell."""
    if _is_blank(value):
        return None, True, True
    if isinstance(value, bool):
        return None, False, False
    if isinstance(value, int):
        return (value, True, False) if value >= 0 else (None, False, False)
    if isinstance(value, float):
        if value.is_integer() and value >= 0:
            return int(value), True, False
        return None, False, False
    if isinstance(value, str):
        stripped = value.strip()
        if re.fullmatch(r'\d+', stripped):
            return int(stripped), True, False
    return None, False, False


def _part_designators(value):
    """Split a designator cell and reject empty segments other than trailing ones."""
    if _is_blank(value):
        return [], []

    segments = [segment.strip() for segment in DESIGNATOR_SEPARATOR.split(str(value))]
    while segments and not segments[-1]:
        segments.pop()
    empty_segments = [
        index for index, segment in enumerate(segments, start=1) if not segment
    ]
    return [segment for segment in segments if segment], empty_segments


def normalize_workbook_rows(rows, source_kind):
    """Add conservative, header-based field hints while preserving source cells."""
    by_sheet = {}
    data_rows = []
    for row in rows:
        by_sheet.setdefault(row['source_sheet'], []).append(row)

    for sheet_rows in by_sheet.values():
        sheet_rows.sort(key=lambda row: row['source_row'])
        header_row = sheet_rows[0]
        labels = {
            column: cell.get('value')
            for column, cell in header_row['source_values'].items()
            if isinstance(cell.get('value'), str) and cell['value'].strip()
        }
        if any(
            label in {'售后状态', '故障原因具体描述', '售后出库', '售后入库'}
            for label in labels.values()
        ):
            profile = 'service_history'
        elif any(
            label
            in {
                '出库产品',
                '流程编号',
                '状态（已出库/未出库）',  # noqa: RUF001 - exact source header
            }
            for label in labels.values()
        ):
            profile = 'shipment_history'
        else:
            profile = source_kind
        header_map = IMPORT_HEADERS.get(profile, {})
        candidates = {
            column: header_map[label.strip()]
            for column, label in labels.items()
            if label.strip() in header_map
        }
        target_counts = {}
        for target in candidates.values():
            target_counts[target] = target_counts.get(target, 0) + 1
        ambiguous = sorted(
            target for target, count in target_counts.items() if count > 1
        )
        column_targets = {
            column: target
            for column, target in candidates.items()
            if target_counts[target] == 1
        }

        if not column_targets:
            for row in sheet_rows:
                row['normalized_values'] = {
                    'row_type': 'data',
                    'mapping_status': 'unrecognized_header',
                    'mapping_profile': profile,
                }
                data_rows.append(row)
            continue

        header_provenance = {
            'source_row': header_row['source_row'],
            'source_range': header_row['source_range'],
            'source_values': header_row['source_values'],
            'column_targets': column_targets,
            'ambiguous_fields': ambiguous,
        }
        for row in sheet_rows[1:]:
            unmapped_columns = sorted(
                column
                for column, cell in row['source_values'].items()
                if column not in column_targets and cell['value'] is not None
            )
            fields = {}
            identifier_warnings = []
            for column, target in column_targets.items():
                cell = row['source_values'][column]
                value = cell['value']
                if value is None:
                    continue
                if (
                    target
                    in {
                        'part_number',
                        'oem_part_number',
                        'serial',
                        'batch_code',
                        'external_reference',
                    }
                    and isinstance(value, (int, float))
                    and not isinstance(value, bool)
                ):
                    number_format = cell.get('number_format', '')
                    if isinstance(value, int) and re.fullmatch(r'0+', number_format):
                        value = str(value).zfill(len(number_format))
                    else:
                        identifier_warnings.append({
                            'field': target,
                            'code': 'identifier_stored_as_number',
                            'source_column': column,
                        })
                fields[target] = value
            row['normalized_values'] = {
                'row_type': 'data',
                'mapping_status': (
                    'partial' if ambiguous or unmapped_columns else 'mapped'
                ),
                'mapping_profile': profile,
                'fields': fields,
                'unmapped_columns': unmapped_columns,
                'ambiguous_fields': ambiguous,
                'identifier_warnings': identifier_warnings,
                'source_headers': header_provenance,
            }
            data_rows.append(row)

    return data_rows


DATE_FIELDS = {'production_date', 'shipped_at', 'returned_at'}
YEAR_PATTERN = re.compile(r'(?<!\d)(?:19|20)\d{2}(?!\d)')
PARTIAL_DATE_PATTERN = re.compile(r'^(?:\d{1,2}[./-]\d{1,2}|\d{1,2}月\d{1,2}日)$')
SERIAL_RANGE_PATTERN = re.compile(
    r'^(?P<start_prefix>[A-Za-z0-9_-]*?)(?P<start>\d+)\s*[-–~至]\s*'  # noqa: RUF001 - accept source range dashes
    r'(?P<end_prefix>[A-Za-z0-9_-]*?)(?P<end>\d+)$'
)


def validate_normalized_import_values(normalized_values):
    """Validate dates and serial ranges without guessing or expanding stock."""
    if not isinstance(normalized_values, dict):
        return []

    fields = normalized_values.get('fields', {})
    if not isinstance(fields, dict):
        return []

    errors = []
    for warning in normalized_values.get('identifier_warnings', []):
        if isinstance(warning, dict):
            errors.append(warning)
    for field in DATE_FIELDS:
        value = fields.get(field)
        if value in (None, ''):
            continue
        if not isinstance(value, str):
            errors.append({'field': field, 'code': 'invalid_date'})
            continue
        value = value.strip()
        if not YEAR_PATTERN.search(value):
            code = (
                'date_missing_year'
                if PARTIAL_DATE_PATTERN.fullmatch(value)
                else 'invalid_date'
            )
            errors.append({'field': field, 'code': code})
            continue
        if not _is_complete_date(value):
            errors.append({'field': field, 'code': 'invalid_date'})

    serial = fields.get('serial')
    quantity = fields.get('quantity')
    if isinstance(serial, str):
        match = SERIAL_RANGE_PATTERN.fullmatch(serial.strip())
        if match:
            start = int(match.group('start'))
            end = int(match.group('end'))
            if match.group('start_prefix') != match.group('end_prefix') or end < start:
                errors.append({'field': 'serial', 'code': 'invalid_serial_range'})
            else:
                expected_quantity = end - start + 1
                try:
                    quantity_matches = Decimal(str(quantity)) == expected_quantity
                except Exception:
                    quantity_matches = False
                errors.append({
                    'field': 'serial',
                    'code': (
                        'serial_range_quantity_mismatch'
                        if not quantity_matches
                        else 'serial_range_requires_manual_expansion'
                    ),
                    'expected_quantity': expected_quantity,
                })

    return errors


def _is_complete_date(value):
    """Accept explicit-year dates only, with a small stable format set."""
    try:
        datetime.fromisoformat(value.replace('Z', '+00:00'))
        return True
    except ValueError:
        pass

    for date_format in ('%Y/%m/%d', '%Y.%m.%d', '%Y年%m月%d日', '%Y%m%d'):
        try:
            datetime.strptime(value, date_format)
            return True
        except ValueError:
            continue
    return False
